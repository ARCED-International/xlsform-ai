"""Template management for XLSForm AI projects."""

import json
import os
import shutil
from datetime import datetime
from pathlib import Path
from typing import List, Optional

from .agents import get_agent

TEMPLATE_VERSION = "0.1.0"
DEFAULT_AGENT = "claude"


class TemplateManager:
    """Manages project templates for XLSForm AI."""

    def __init__(self, template_dir: Optional[Path] = None):
        """Initialize template manager.

        Args:
            template_dir: Directory containing templates. Defaults to templates/ in package.
        """
        if template_dir is None:
            # Templates should be inside the xlsform_ai package
            import importlib.resources as resources

            try:
                # Try to find templates in the installed package
                with resources.files("xlsform_ai") as pkg_files:
                    templates_path = pkg_files / "templates"
                    if templates_path.is_dir():
                        template_dir = Path(str(templates_path))
            except Exception:
                # Fallback to relative path for development
                template_dir = Path(__file__).parent / "templates"

        self.template_dir = Path(template_dir)
        self.base_template = self.template_dir / "base"

    def init_project(
        self,
        project_path: Path,
        agents: List[str] = None,
        overwrite: bool = False,
    ) -> bool:
        """Initialize a new XLSForm AI project.

        Args:
            project_path: Path where project should be initialized
            agents: List of agents to configure for (default: [claude])
            overwrite: Whether to overwrite existing files

        Returns:
            True if successful, False otherwise
        """
        # Default to claude if no agents specified
        if agents is None:
            agents = [DEFAULT_AGENT]

        # Validate all agents
        for agent in agents:
            agent_config = get_agent(agent)
            if not agent_config:
                print(f"Error: Agent '{agent}' is not supported.")
                from .agents import get_supported_agents
                print(f"Supported agents: {', '.join(get_supported_agents())}")
                return False

        # Check if template exists
        if not self.base_template.exists():
            print(f"Error: Template not found at {self.base_template}")
            return False

        # Create project directory if it doesn't exist
        project_path = Path(project_path)
        project_path.mkdir(parents=True, exist_ok=True)

        # Copy template files
        success = self._copy_template_files(
            project_path,
            agents,
            overwrite
        )

        return success

    def _copy_template_files(
        self,
        project_path: Path,
        agents: List[str],
        overwrite: bool,
    ) -> bool:
        """Copy template files to project directory.

        Args:
            project_path: Target project directory
            agents: List of agent names to configure for
            overwrite: Whether to overwrite existing files

        IMPORTANT: Never overwrite survey file and activity log files during re-init,
        even if overwrite=True. These files contain user data.

        Returns:
            True if successful
        """
        try:
            # Detect existing survey file (from config or default)
            try:
                # Import config module from scripts
                import sys
                scripts_dir = self.base_template / "scripts"
                if str(scripts_dir) not in sys.path:
                    sys.path.insert(0, str(scripts_dir))
                from config import ProjectConfig, DEFAULT_CONFIG

                # Check if config exists in project
                config_file = project_path / "xlsform-ai.json"
                if config_file.exists():
                    config = ProjectConfig(project_path)
                    survey_file_name = config.get_xlsform_file()
                else:
                    survey_file_name = "survey.xlsx"
            except Exception:
                survey_file_name = "survey.xlsx"

            survey_xlsx = project_path / survey_file_name

            # Find existing activity log files
            existing_activity_logs = list(project_path.glob("activity_log_*.html"))

            # Copy .claude directory
            claude_src = self.base_template / ".claude"
            claude_dst = project_path / ".claude"

            if claude_dst.exists():
                if overwrite:
                    shutil.rmtree(claude_dst)
                else:
                    print(f"Note: .claude directory already exists, skipping...")

            if not claude_dst.exists() or overwrite:
                shutil.copytree(claude_src, claude_dst)
                print(f"✓ Created .claude directory")

            # For each agent, copy agent-specific files if they exist
            for agent in agents:
                agent_template_dir = self.template_dir / agent

                if agent_template_dir.exists():
                    self._merge_agent_config(claude_dst, agent_template_dir)
                    print(f"✓ Added {agent} configuration")

            # Copy scripts directory
            scripts_src = self.base_template / "scripts"
            scripts_dst = project_path / "scripts"

            if scripts_dst.exists():
                if overwrite:
                    shutil.rmtree(scripts_dst)
                else:
                    print(f"Note: scripts directory already exists, skipping...")

            if not scripts_dst.exists() or overwrite:
                shutil.copytree(scripts_src, scripts_dst)
                print(f"✓ Created scripts directory")

            # Copy package.json if it exists
            package_json_src = self.base_template / "package.json"
            package_json_dst = project_path / "package.json"

            if package_json_src.exists():
                if package_json_dst.exists() and not overwrite:
                    print(f"Note: package.json already exists, skipping...")
                else:
                    shutil.copy2(package_json_src, package_json_dst)
                    print(f"✓ Created package.json")

            # Create configuration file
            config_file = project_path / "xlsform-ai.json"
            if not config_file.exists() or overwrite:
                try:
                    config_data = DEFAULT_CONFIG.copy()
                    config_data["project_name"] = project_path.name
                    config_data["created"] = datetime.now().isoformat()

                    with open(config_file, 'w', encoding='utf-8') as f:
                        json.dump(config_data, f, indent=2)

                    print(f"✓ Created xlsform-ai.json configuration")
                except Exception as e:
                    print(f"Note: Could not create config file: {e}")

            # Create example XLSForm file - NEVER overwrite if it exists
            if not survey_xlsx.exists():
                template_xlsx = self.base_template / ".claude" / "skills" / "xlsform-core" / "assets" / "xlsform-template.xlsx"
                if template_xlsx.exists():
                    shutil.copy2(template_xlsx, survey_xlsx)
                    # Apply freeze panes
                    try:
                        import openpyxl
                        wb = openpyxl.load_workbook(survey_xlsx)
                        for sheet in wb.worksheets:
                            for row_idx in range(1, min(10, sheet.max_row + 1)):
                                if sheet.cell(row_idx, 1).value and str(sheet.cell(row_idx, 1).value).strip().lower() == "type":
                                    sheet.freeze_panes = f"A{row_idx + 1}"
                                    break
                        wb.save(survey_xlsx)
                    except Exception as e:
                        print(f"Note: Could not apply freeze panes: {e}")
                else:
                    survey_xlsx.touch()
                print(f"✓ Created {survey_file_name}")
            else:
                print(f"Note: {survey_file_name} already exists, preserving data...")

            # NEVER delete activity log files during re-init
            if existing_activity_logs:
                print(f"Note: Preserving {len(existing_activity_logs)} existing activity log file(s)")

            return True

        except Exception as e:
            print(f"Error copying template files: {e}")
            return False

    def _merge_agent_config(
        self,
        target_dir: Path,
        agent_template_dir: Path,
    ):
        """Merge agent-specific configuration into target directory.

        Args:
            target_dir: Target .claude directory
            agent_template_dir: Agent-specific template directory
        """
        # Copy agent-specific files, overwriting base files if they exist
        for item in agent_template_dir.iterdir():
            target = target_dir / item.name

            if item.is_dir():
                if target.exists():
                    shutil.rmtree(target)
                shutil.copytree(item, target)
            else:
                shutil.copy2(item, target)

    def get_template_path(self) -> Path:
        """Get the path to the base template.

        Returns:
            Path to base template directory
        """
        return self.base_template
